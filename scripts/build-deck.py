#!/usr/bin/env python3
"""Build the delta-git presentation deck.

python-pptx -> docs/hackathon/video/slides/delta-git-deck.pptx
Companion teleprompter script -> docs/hackathon/video/slides/narration.md

16:9, dark theme matching the product, gold/green Widespread accent.
Each slide carries speaker notes; the same copy drives narration.md.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image
import os

SHOTS = "docs/hackathon/video/slides/shots"
OUT_DIR = "docs/hackathon/video/slides"

BG = RGBColor(0x0D, 0x11, 0x17)
PANEL = RGBColor(0x16, 0x1B, 0x22)
FG = RGBColor(0xE6, 0xED, 0xF3)
DIM = RGBColor(0x8B, 0x94, 0x9E)
GOLD = RGBColor(0xD2, 0x99, 0x22)
GREEN = RGBColor(0x3F, 0xB9, 0x50)
BLUE = RGBColor(0x58, 0xA6, 0xFF)

# (kind, title, caption, image, notes)
# kind: "title" | "text" | "shot"
SLIDES = [
    (
        "title",
        "delta-git",
        "The forge where pushes never get rejected",
        None,
        "Hi — this is delta-git, a git forge rebuilt for the agentic era on "
        "Cloudflare Workers and Artifacts. The pitch in one line: on this "
        "forge, a push is never rejected. Every competing change lands, gets "
        "catalogued, and gets merged — humans and agents alike. It's live at "
        "code.widespread.fyi. Let me walk you through what was built.",
    ),
    (
        "text",
        "The problem",
        "GitHub serializes writes.\n\n"
        "• Push a divergent commit → remote rejected\n"
        "• Fetch, rebase, force-push — or lose the work\n"
        "• Fine for humans at human speed\n"
        "• Broken for hundreds of agents pushing concurrently\n\n"
        "The fix isn't a better merge queue.\n"
        "It's treating divergence as data.",
        None,
        "The problem is the write model. On GitHub, if two actors push to the "
        "same branch, the second one gets rejected — fetch, rebase, force-push, "
        "or lose the work. Humans tolerate that. Agents don't: at scale, "
        "concurrent writes are the norm, and a retry loop just moves the "
        "contention. We didn't build a faster merge queue. We made divergence "
        "itself a first-class object.",
    ),
    (
        "shot",
        "Idea-first front door",
        "Non-coders post ideas — agents claim, spec, patch, and land them behind quorum review.",
        "landing-anon.png",
        "This is the front door, and it's deliberately not a repository list. "
        "The first thing a visitor sees is the idea composer: 'What should the "
        "agents build?' Below it, a Connect-an-agent panel — the skill doc, "
        "the MCP endpoint, the OpenAPI spec — because the primary audience is "
        "mixed: non-coders bring ideas, agents build them, humans verify. No "
        "sign-in wall to browse.",
    ),
    (
        "shot",
        "atproto-native identity",
        "Sign in with a DID — did:plc, did:key, or a Bluesky handle. Keys never leave the client.",
        "signin.png",
        "Identity is atproto-native. You can sign in with a Bluesky handle "
        "through client-side OAuth — the tokens never transit our worker — or "
        "with a bare did:key by signing a challenge. Server never sees a "
        "private key, sessions are DPoP-bound, and agent keys are bound and "
        "revoked through the same DID machinery.",
    ),
    (
        "shot",
        "Public repos don't ride our infra",
        "Explore — public repos mirror out to Tangled and Radicle over SSH.",
        "explore.png",
        "Explore lists public repositories. Public work doesn't just live on "
        "our infrastructure — federation mirrors pushes out to Tangled and "
        "Radicle over SSH from inside a Cloudflare Container, signed by a "
        "platform mirror identity. Your public repo survives us.",
    ),
    (
        "shot",
        "A normal repo — mostly",
        "Code, issues, PRs, CI — plus a banner for the delta ref that just landed.",
        "repo-summary.png",
        "The day-to-day is deliberately boring: code browser, issues, pull "
        "requests, commits — the GitHub-shaped surface everyone knows. But "
        "look at the banner: 'delta/main-1e7457e2d999 had recent pushes.' A "
        "competing commit just landed on a delta ref, and the UI offers a "
        "compare-and-merge path instead of an error.",
    ),
    (
        "shot",
        "Two pushes raced. Neither lost.",
        "refs/delta/main-1e7457e2d999 — quarantined, catalogued, waiting for adjudication.",
        "repo-intents.png",
        "Here's the whole thesis in one screenshot. Two pushes hit main at the "
        "same time — real receive-pack calls, real concurrency. The first "
        "landed on main. The second was not rejected: it was rewritten to "
        "refs/delta/main-1e7457e2d999 and minted this open merge intent — "
        "target, delta ref, actor, status, conflicts. The work exists. It's "
        "catalogued. It waits for adjudication, not a retry.",
    ),
    (
        "text",
        "How never-reject works",
        "• Receives stage in parallel — an 8-lease pool per repo DO\n"
        "• Finalize is a single-writer critical section\n"
        "• Ref moved under you → delta ref + merge intent\n"
        "• Quorum adjudicates: auto-merge, votes, or Clef AI seat\n"
        "• Policy gates quarantine the same way — never reject\n\n"
        "Measured: 50 concurrent pushes → 0 rejections, 9 intents, ~1.4s",
        None,
        "The mechanism: each repository is a Durable Object. Receive packs "
        "stage in parallel through an eight-lease pool, but finalization is a "
        "single-writer critical section — that's what makes ordering honest. "
        "If the ref moved while your pack was in flight, the push is "
        "re-targeted to a delta ref and a merge intent is minted. A quorum "
        "then adjudicates — clean cases auto-merge, contested ones get "
        "reputation-weighted votes plus a structured Clef AI decision seat. "
        "And when a policy or freeze window blocks a push, it quarantines the "
        "same way — the contract is absolute. We measured it: fifty "
        "concurrent pushes, zero rejections, nine merge intents, about 1.4 "
        "seconds of wall time.",
    ),
    (
        "shot",
        "Arena — agents compete on git history",
        "Real matches, resolved: isolated forks → blind judging → winner merges to canonical.",
        "arena-global.png",
        "The arena reuses the same machinery as a feature: a match gives each "
        "entrant an isolated Artifacts workspace fork, they push real work, "
        "and entries are scored by automatic signals plus blind community "
        "votes — voters see the diff, not the author. These are real resolved "
        "matches from production. The winner merges to canonical with the "
        "same adjudication path as any other merge intent.",
    ),
    (
        "shot",
        "One reputation currency",
        "Humans and agents on the same leaderboard — vouches, epochs, rep-gated actions.",
        "agents.png",
        "Reputation is one currency across humans and agents — arena bots "
        "here earned rep the same way a person would: vouches, Coordinape-"
        "style epoch allocations, and adjudication outcomes. And rep is not "
        "decorative — it gates vote weight, intent claiming, and sensitive "
        "operations.",
    ),
    (
        "shot",
        "Dev environments in real containers",
        "Workspaces run in Cloudflare Containers — terminal, file editor, authenticated preview.",
        "dev-environments.png",
        "Dev environments aren't a status field — starting one boots a real "
        "workspace in a Cloudflare Container, with a terminal, a file editor, "
        "and an authenticated preview URL. Egress is allowlisted per job type "
        "and the mirror path fails closed.",
    ),
    (
        "shot",
        "CI runs where the code lives",
        "The built-in delegate runner claims executions and runs them in a ForgeSandbox.",
        "repo-pipelines.png",
        "Pipelines are dogfooded — this repo's own .delta/workflows/ci.yml "
        "runs here. A built-in delegate runner claims executions off the "
        "ledger, stages the worktree into a ForgeSandbox container, and "
        "streams logs back. Same protocol external runners use — we just ship "
        "one in the box.",
    ),
    (
        "shot",
        "Chaos as a product feature",
        "Endpoint chaos + forge self-chaos — measured lease contention, not vibes.",
        "chaos.png",
        "Chaos isn't a marketing word — the chaos runner mints a real PAT and "
        "floods real endpoints, and the forge self-chaos probe injects "
        "contention directly into the receive lease pool and measures it. The "
        "run you see reports actual grant/busy counts.",
    ),
    (
        "shot",
        "GitOps — pushes reconcile deploys",
        "Environments as deploy targets; drift stays open until the deploy actually lands.",
        "gitops.png",
        "GitOps targets watch tracked branches — a push enqueues a real "
        "deploy to the named environment, and the drift ledger only converges "
        "when the deploy actually lands. Honest by construction.",
    ),
    (
        "shot",
        "Databases without the dashboard",
        "Provision, migrate, and query real D1 from the repo UI.",
        "databases.png",
        "Databases are real too — the module provisions actual Cloudflare D1 "
        "databases through the account API, runs migrations, and gives you a "
        "SQL console. No mock lifecycle.",
    ),
    (
        "shot",
        "Supply chain, signed end to end",
        "gitleaks, trivy, osv-scanner, syft in-container — SBOMs and DSSE-signed bundles.",
        "supply-chain.png",
        "Security testing runs pinned versions of gitleaks, trivy, "
        "osv-scanner and syft inside the ForgeSandbox — checksum-verified at "
        "image build. SBOMs are generated for every build, and provenance "
        "bundles are DSSE-signed so an artifact's history is verifiable "
        "off-platform.",
    ),
    (
        "shot",
        "Policy that quarantines, never rejects",
        "Freeze windows and overrides hold pushes on delta refs — work is never lost.",
        "overrides.png",
        "This is the detail that makes the model coherent: even a freeze "
        "window doesn't reject a push. The push is quarantined onto a delta "
        "ref with a merge intent, and an explicit override — recorded, "
        "scoped, auditable — releases it. Policy becomes scheduling, not a "
        "wall.",
    ),
    (
        "shot",
        "Agent-native by construction",
        "OpenAPI, MCP + WebMCP tools, skill.md, llms.txt, poppy.json — same end states as the UI.",
        "api-docs.png",
        "Everything you saw has an agent surface. The REST API is "
        "OpenAPI-described — this is the Scalar reference. MCP and WebMCP "
        "expose tools directly. skill.md and llms.txt teach agents the "
        "conventions, and .well-known/poppy.json declares the forge to "
        "personal agents. We verified parity mechanically: a hundred and "
        "eleven endpoints, every UI end state reachable by REST.",
    ),
    (
        "text",
        "All of it on Cloudflare",
        "Workers — API, git protocol, SPA\n"
        "Durable Objects — per-repo authority, adjudicators, firehose\n"
        "Artifacts — canonical repos, workspace forks, push events\n"
        "Containers — ForgeSandbox jobs, dev environments\n"
        "D1 · R2 · KV · Queues · Vectorize\n"
        "Analytics Engine · Workers AI (llama + Clef seat)",
        None,
        "And it is Cloudflare all the way down. Workers serve the API, the "
        "git smart-HTTP protocol, and the SPA. Durable Objects hold per-repo "
        "authority — refs, receive leases, merge intents — plus adjudicator, "
        "repo-agent, and firehose agents. Cloudflare Artifacts hosts "
        "canonical repos and workspace forks. Cloudflare Containers run "
        "sandboxes and dev environments. D1, R2, KV, Queues, Vectorize, "
        "Analytics Engine, Workers AI — no other vendor in the request path.",
    ),
    (
        "text",
        "Measured, not promised",
        "• 50 concurrent pushes — 0 rejections, 9 merge intents, ~1.4s\n"
        "• 771 worker integration tests · 107 unit tests\n"
        "• 111-endpoint agent-parity suite — UI ≡ REST\n"
        "• 47/47 Playwright checks on the live deploy\n"
        "• MIT · github.com/TerexitariusStomp/delta-git\n\n"
        "code.widespread.fyi",
        None,
        "To close: none of this is a mock. The load profile is in the repo — "
        "fifty concurrent pushes, zero rejections. Seven hundred and "
        "seventy-one worker integration tests, a hundred and eleven agent-"
        "parity endpoints, forty-seven of forty-seven Playwright checks "
        "against the live deployment. It's MIT-licensed, it's open, and it's "
        "running right now at code.widespread.fyi. Thanks — questions "
        "welcome.",
    ),
]

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def add_bg(slide):
    from pptx.enum.shapes import MSO_SHAPE

    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    r.fill.solid()
    r.fill.fore_color.rgb = BG
    r.line.fill.background()
    r.shadow.inherit = False
    return r


def add_text(slide, x, y, w, h, runs, size=18, color=FG, bold=False, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    lines = runs if isinstance(runs, list) else runs.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.bold = bold
    return tb


def accent_bar(slide, x, y, w=Inches(0.9)):
    from pptx.enum.shapes import MSO_SHAPE

    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, Emu(int(Inches(0.045))))
    bar.fill.solid()
    bar.fill.fore_color.rgb = GOLD
    bar.line.fill.background()
    bar.shadow.inherit = False


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


for kind, title, caption, image, notes in SLIDES:
    slide = prs.slides.add_slide(BLANK)
    add_bg(slide)

    if kind == "title":
        add_text(slide, Inches(1.2), Inches(2.3), Inches(10.9), Inches(1.4), title, 66, FG, True)
        accent_bar(slide, Inches(1.25), Inches(3.7))
        add_text(slide, Inches(1.2), Inches(4.0), Inches(10.9), Inches(0.7), caption, 28, GOLD, True)
        add_text(
            slide,
            Inches(1.2),
            Inches(5.0),
            Inches(10.9),
            Inches(1.2),
            "code.widespread.fyi  ·  Cloudflare Workers + Artifacts  ·  MIT",
            18,
            DIM,
        )
    elif kind == "text":
        add_text(slide, Inches(0.9), Inches(0.55), Inches(11.5), Inches(1.0), title, 40, FG, True)
        accent_bar(slide, Inches(0.95), Inches(1.45))
        add_text(slide, Inches(0.9), Inches(1.9), Inches(11.5), Inches(5.2), caption, 22, FG)
    else:  # shot
        add_text(slide, Inches(0.7), Inches(0.35), Inches(11.9), Inches(0.8), title, 32, FG, True)
        accent_bar(slide, Inches(0.75), Inches(1.05))
        path = os.path.join(SHOTS, image)
        iw, ih = Image.open(path).size
        # fit within 11.2 x 5.3 in, centered
        max_w, max_h = Inches(11.2), Inches(5.3)
        scale = min(max_w / iw, max_h / ih)
        w, h = int(iw * scale), int(ih * scale)
        left = int((prs.slide_width - w) / 2)
        top = int(Inches(1.25))
        pic = slide.shapes.add_picture(path, left, top, width=w, height=h)
        pic.line.color.rgb = PANEL
        pic.line.width = Pt(1.5)
        add_text(
            slide,
            Inches(0.7),
            Inches(6.72),
            Inches(11.9),
            Inches(0.6),
            caption,
            15,
            DIM,
            align=PP_ALIGN.CENTER,
        )

    add_notes(slide, notes)

prs.save(f"{OUT_DIR}/delta-git-deck.pptx")

# Teleprompter script — same copy, timed.
with open(f"{OUT_DIR}/narration.md", "w") as f:
    f.write("# delta-git deck narration (~7 min)\n\n")
    f.write("Slide-by-slide speaker notes, identical to the .pptx notes panes.\n\n")
    for i, (_, title, _cap, _img, notes) in enumerate(SLIDES, 1):
        f.write(f"## {i}. {title}\n\n{notes}\n\n")

print(f"deck: {OUT_DIR}/delta-git-deck.pptx ({len(SLIDES)} slides)")
print(f"notes: {OUT_DIR}/narration.md")
